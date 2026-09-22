"""How the workbook reads when someone scrolls it.

Checked by loading the written file with a different library than the one that
wrote it, because the assertion that matters is what Excel sees, not what
xlsxwriter was asked for.
"""

from __future__ import annotations

import openpyxl
import pandas as pd
import pytest

from abductionbench.core.reporting import _write_sheet

SHEETS = ("Summary_Long", "Summary", "S_abd", "Metrics")


@pytest.fixture
def workbook(tmp_path):
    long_df = pd.DataFrame(
        {
            "dataset_id": ["abd", "aer"],
            "model_id": ["m1", "m2"],
            "prompt_mode": ["io", "cot"],
            "selection_mode": ["n-a", "MCS"],
            "value": [0.5, 0.25],
            "metric.reasoning_observation_total": [3, 4],
        }
    )
    pivot = pd.DataFrame(
        {"m1": [0.5], "m2": [0.25]}, index=pd.Index(["abd"], name="dataset_id")
    )
    path = tmp_path / "abductionbench_results.xlsx"
    with pd.ExcelWriter(path, engine="xlsxwriter") as writer:
        _write_sheet(writer, long_df, "Summary_Long", freeze_columns=4)
        _write_sheet(writer, pivot, "Summary", index=True)
        _write_sheet(writer, long_df, "S_abd")
        _write_sheet(writer, long_df, "Metrics")
    return openpyxl.load_workbook(path)


@pytest.mark.parametrize("sheet", SHEETS)
def test_every_body_cell_is_centred_both_ways(workbook, sheet):
    worksheet = workbook[sheet]
    seen = {
        (worksheet.cell(row=row, column=col).alignment.horizontal,
         worksheet.cell(row=row, column=col).alignment.vertical)
        for row in range(2, worksheet.max_row + 1)
        for col in range(1, worksheet.max_column + 1)
    }
    assert seen == {("center", "center")}, f"{sheet}: {seen}"


@pytest.mark.parametrize("sheet", SHEETS)
def test_every_header_cell_is_centred_both_ways(workbook, sheet):
    worksheet = workbook[sheet]
    for col in range(1, worksheet.max_column + 1):
        alignment = worksheet.cell(row=1, column=col).alignment
        assert (alignment.horizontal, alignment.vertical) == ("center", "center")


def test_summary_long_freezes_the_four_identity_columns(workbook):
    """dataset_id, model_id, prompt_mode, selection_mode.

    Scrolled sideways to the reasoning columns, a row without them says
    nothing about which task it is.
    """
    assert workbook["Summary_Long"].freeze_panes == "E2"


def test_summary_freezes_its_first_column(workbook):
    """The dataset label, so a wide dataset x model matrix stays readable."""
    assert workbook["Summary"].freeze_panes == "B2"


def test_a_plain_sheet_still_freezes_only_the_header(workbook):
    """The default is unchanged: no identity columns, nothing to pin."""
    assert workbook["S_abd"].freeze_panes == "A2"


def test_freezing_never_exceeds_the_sheets_width(tmp_path):
    """A narrow frame asked to freeze four columns must not freeze past its edge."""
    narrow = pd.DataFrame({"only": [1, 2]})
    path = tmp_path / "narrow.xlsx"
    with pd.ExcelWriter(path, engine="xlsxwriter") as writer:
        _write_sheet(writer, narrow, "Summary_Long", freeze_columns=4)
    assert openpyxl.load_workbook(path)["Summary_Long"].freeze_panes == "B2"
