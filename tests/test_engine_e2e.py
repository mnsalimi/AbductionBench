"""End-to-end engine tests over real HTTP against a fake vLLM server.

These cover the behaviours that make the engine trustworthy in production:
native batching with the configured group size, per-sample checkpointing and
resume, retry on transient failures, bisection of a batch poisoned by one bad
sample, oversize-sample replacement, single-call fallback, and reporting.
"""

from __future__ import annotations

import asyncio
import json
from pathlib import Path

import orjson
import pytest

from abductionbench.core.config import load_run_config
from abductionbench.core.engine import EvaluationEngine
from abductionbench.core.reporting import write_reports


def _run(config_path: Path, **kwargs):
    config = load_run_config(config_path)
    engine = EvaluationEngine(config, **kwargs)
    return asyncio.run(engine.run()), engine


def _records(task_dir: Path) -> list[dict]:
    path = task_dir / "records.jsonl"
    return [orjson.loads(line) for line in path.read_bytes().splitlines() if line.strip()]


def test_batched_run_produces_records_and_metrics(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=8, sample_size=8)]
    )
    result, _ = _run(config_path)

    assert len(result.tasks) == 1
    task = result.tasks[0]
    assert task.n_planned == 8
    assert task.n_scored == 8
    assert task.n_error == 0
    assert task.metrics["coverage"] == 1.0
    # The first batch call is the engine's startup probe (2 conversations);
    # then group_size=4 over 8 samples -> exactly two batch calls, run server-side.
    assert fake_server.state.batch_calls == [2, 4, 4]
    assert fake_server.state.single_calls == 0

    records = _records(task.output_dir)
    assert len(records) == 8
    assert all(record["status"] == "ok" for record in records)
    assert all(record["response"]["batch_size"] == 4 for record in records)
    # Batch usage is an aggregate; the record says so rather than pretending
    # it is per-sample.
    assert all(record["response"]["usage_is_batch_aggregate"] for record in records)
    assert (task.output_dir / "metrics.json").exists()
    assert (task.output_dir / "checkpoint.json").exists()
    assert list((task.output_dir / "raw").glob("*.json"))


def test_per_model_group_size_is_honoured(fake_server, write_run_config, fake_dataset):
    models = [
        {
            "id": "m8",
            "model_name": "test/model",
            "endpoint": {
                "base_url": fake_server.base_url,
                "api_key": "test-key",
                "batch": {"enabled": True, "group_size": 8},
            },
            "sampling": {"max_tokens_default": 64},
            "limits": {"max_parallel_batches": 1},
        },
        {
            "id": "m3",
            "model_name": "test/model",
            "endpoint": {
                "base_url": fake_server.base_url,
                "api_key": "test-key",
                "batch": {"enabled": True, "group_size": 3},
            },
            "sampling": {"max_tokens_default": 64},
            "limits": {"max_parallel_batches": 1},
        },
    ]
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=6, sample_size=6)],
        models=models,
        engine={"concurrency": {"max_parallel_tasks": 1}},
    )
    result, _ = _run(config_path)
    assert len(result.tasks) == 2
    # 6 samples: model m8 -> one call of 6; model m3 -> two calls of 3.
    # The verification probe also issues one 2-conversation call per model.
    sizes = sorted(fake_server.state.batch_calls)
    assert sizes.count(2) == 2  # probes
    assert 6 in sizes and sizes.count(3) == 2


def test_resume_skips_completed_samples(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=8, sample_size=8)]
    )
    result, engine = _run(config_path)
    calls_first = len(fake_server.state.batch_calls)

    # Re-run into the same directory: every sample is checkpointed already.
    config = load_run_config(config_path)
    engine2 = EvaluationEngine(config, run_id=result.run_id, run_dir=result.run_dir)
    result2 = asyncio.run(engine2.run())
    task = result2.tasks[0]
    assert task.n_reused == 8
    assert task.n_scored == 8
    # Only the endpoint verification probe was issued the second time.
    assert len(fake_server.state.batch_calls) == calls_first + 1


def test_resume_is_invalidated_by_a_mode_change(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=4, sample_size=4)]
    )
    result, _ = _run(config_path)

    changed = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=4, sample_size=4)],
        modes={"prompt_modes": ["cot"]},
        name="test-run-2",
    )
    config = load_run_config(changed)
    engine = EvaluationEngine(config, run_id=result.run_id, run_dir=result.run_dir)
    result2 = asyncio.run(engine.run())
    # A different prompt mode is a different run version: its own task
    # directory, and nothing reused from the io run.
    assert result2.tasks[0].n_reused == 0
    assert result2.tasks[0].identity.prompt_mode == "cot"
    assert result2.tasks[0].identity.template_mode == "cot|n/a|generation|static"


def test_transient_failures_are_retried(fake_server, write_run_config, fake_dataset):
    fake_server.state.fail_next = [503, 429]  # both consumed by the first calls
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=4, sample_size=4)]
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    assert task.n_scored == 4
    assert task.n_error == 0
    events = [
        json.loads(line)
        for line in (result.run_dir / "events.jsonl").read_text().splitlines()
        if line.strip()
    ]
    assert any(event["event"] == "batch_attempt_failed" for event in events)


def test_permanent_failure_is_recorded_per_sample(fake_server, write_run_config, fake_dataset):
    fake_server.state.fail_next = [503] * 40  # exceeds max_attempts for every call
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=4, sample_size=4)]
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    assert task.n_error == 4
    assert task.n_scored == 0
    assert task.metrics["coverage"] == 0.0
    statuses = {record["status"] for record in _records(task.output_dir)}
    assert statuses == {"error"}


def test_poison_sample_is_isolated_by_bisection(fake_server, write_run_config, fake_dataset):
    """One bad conversation fails the whole batch; the rest must still land."""
    fake_server.state.poison_marker = "POISON"
    config_path = write_run_config(
        base_url=fake_server.base_url,
        # sample index 0 carries the marker (poison_every=8 -> only s0000)
        datasets=[fake_dataset("fake", n=8, sample_size=8, poison_every=8)],
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    records = {record["sample_id"]: record for record in _records(task.output_dir)}
    assert len(records) == 8
    # The offending sample is skipped (context-length class), everything else scored.
    assert records["s0000"]["status"] == "skipped"
    assert sum(1 for r in records.values() if r["status"] == "ok") == 7
    assert task.checkpoint.bisections >= 1


def test_oversize_samples_are_replaced_not_dropped(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[
            fake_dataset("fake", n=6, sample_size=6, long_every=3, long_chars=20000, replacements=10)
        ],
        engine={"limits": {"input_token_budget": 2000, "on_oversize": "resample"}},
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    # Oversize samples (every 3rd, including a drawn replacement) were replaced
    # rather than dropped, so the evaluated set keeps its configured size.
    assert task.n_planned == 6
    assert task.n_scored == 6
    doc = task.documentation
    assert doc.statistics["oversize_dropped[io_n-a_static]"] == 3
    assert doc.statistics["replacements_used[io_n-a_static]"] == 3
    assert doc.statistics["prompts_kept[io_n-a_static]"] == 6


def test_oversize_skip_policy_shrinks_the_set(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=6, sample_size=6, long_every=3, long_chars=20000)],
        engine={"limits": {"input_token_budget": 2000, "on_oversize": "skip"}},
    )
    result, _ = _run(config_path)
    assert result.tasks[0].n_planned == 4


def test_dataset_mostly_oversize_is_reported_unusable(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[
            fake_dataset("fake", n=4, sample_size=4, long_every=1, long_chars=20000, replacements=0)
        ],
        engine={"limits": {"input_token_budget": 1000, "on_oversize": "resample"}},
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    assert task.failure and "input budget" in task.failure


def test_empty_response_is_its_own_status(fake_server, write_run_config, fake_dataset):
    """A reasoning model can burn max_tokens and return content: null."""
    fake_server.state.empty_marker = "observation number 1"
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=4, sample_size=4)]
    )
    result, _ = _run(config_path)
    records = {r["sample_id"]: r for r in _records(result.tasks[0].output_dir)}
    assert records["s0001"]["status"] == "empty"
    assert records["s0001"]["parse_ok"] is False
    assert result.tasks[0].metrics["empty_response_rate"] == pytest.approx(0.25)


def test_fallback_to_single_calls_when_batch_route_missing(
    fake_server, write_run_config, fake_dataset
):
    fake_server.state.batch_enabled = False
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=4, sample_size=4)]
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    assert task.n_scored == 4
    assert fake_server.state.single_calls == 4
    assert task.diagnostics["batch_mode"] is False


def test_batch_route_missing_without_fallback_aborts(fake_server, write_run_config, fake_dataset):
    fake_server.state.batch_enabled = False
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=2, sample_size=2)],
        engine={"batching": {"fallback_to_single": False}},
    )
    from abductionbench.core.errors import AbenchError

    with pytest.raises(AbenchError, match="batch endpoint unusable"):
        _run(config_path)


def test_bad_api_key_fails_fast(fake_server, write_run_config, fake_dataset):
    fake_server.state.api_key = "correct-key"
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=2, sample_size=2)]
    )
    from abductionbench.core.errors import AbenchError

    with pytest.raises(AbenchError, match="authentication failed"):
        _run(config_path)


def test_skipped_dataset_does_not_stop_the_run(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[
            fake_dataset("broken", skip="source repository is unreachable"),
            fake_dataset("fine", n=4, sample_size=4),
        ],
    )
    result, _ = _run(config_path)
    assert [entry["dataset_id"] for entry in result.skipped_datasets] == ["broken"]
    assert [task.identity.dataset_id for task in result.tasks] == ["fine"]


def test_prompt_modes_create_parallel_tasks(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=4, sample_size=4)],
        modes={"prompt_modes": ["io", "cot"]},
    )
    result, _ = _run(config_path)
    assert sorted(task.identity.prompt_mode for task in result.tasks) == ["cot", "io"]
    # Both modes evaluated the same samples, and each is its own run version.
    assert {task.n_planned for task in result.tasks} == {4}
    assert len({task.identity.template_mode for task in result.tasks}) == 2


def test_self_consistency_votes_are_reduced_to_one_score(
    fake_server, write_run_config, fake_dataset
):
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=4, sample_size=4)],
        modes={"prompt_modes": ["self-consistency"], "self_consistency_n": 3},
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    # Four items asked three times each: twelve requests, four scored answers.
    assert task.n_planned == 12
    assert task.n_scored == 4
    assert task.metrics["self_consistency_agreement"] == 1.0


def test_dry_run_calls_nothing(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=4, sample_size=4)]
    )
    result, _ = _run(config_path, dry_run=True)
    assert fake_server.state.requests == 0
    assert result.tasks[0].n_planned == 4


def test_reports_are_written(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("a", n=4, sample_size=4), fake_dataset("b", n=4, sample_size=4)],
    )
    result, _ = _run(config_path)
    written = write_reports(result)
    excel = written["excel"]
    assert excel.exists() and excel.stat().st_size > 0
    assert (result.run_dir / "RUN_REPORT.md").exists()
    assert (written["summary_csv"]).exists()
    for task in result.tasks:
        assert (task.output_dir / "run_documentation.md").exists()

    import pandas as pd

    sheets = pd.read_excel(excel, sheet_name=None)
    assert {"Summary", "Metrics", "Tasks", "Datasets", "Skipped", "Models"} <= set(sheets)
    assert any(name.startswith("S_") for name in sheets)


def test_report_command_rebuilds_from_records(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=4, sample_size=4)]
    )
    result, _ = _run(config_path)
    from abductionbench.core.rebuild import rebuild_run_result

    rebuilt = rebuild_run_result(result.run_dir)
    assert len(rebuilt.tasks) == 1
    assert rebuilt.tasks[0].n_scored == 4
    assert rebuilt.tasks[0].metrics["coverage"] == 1.0


def test_exhausted_transient_retries_bisect_the_batch(fake_server, write_run_config, fake_dataset):
    """A batch too slow/flaky as a whole is halved rather than failed wholesale.

    Four injected 503s exhaust the retry budget of the first (size-4) call and
    of its first half, leaving the second half to succeed -- so half the samples
    survive a failure that would otherwise have lost all four.
    """
    fake_server.state.fail_next = [503] * 4
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=4, sample_size=4)],
        engine={"batching": {"bisect_on_timeout": True}},
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    assert task.checkpoint.bisections >= 1
    # Samples are salvaged by the split instead of the whole batch being lost.
    assert task.n_scored >= 2
    assert task.n_scored + task.n_error == 4
    assert task.diagnostics["batch_mode"] is True


def test_timeout_bisection_can_be_disabled(fake_server, write_run_config, fake_dataset):
    fake_server.state.fail_next = [503] * 4
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=4, sample_size=4)],
        engine={"batching": {"bisect_on_timeout": False}},
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    assert task.checkpoint.bisections == 0
    assert task.n_error == 4


def test_transient_probe_failure_keeps_batch_mode(fake_server, write_run_config, fake_dataset):
    """A momentary failure of the startup probe must not downgrade the run."""
    fake_server.state.fail_next = [503]  # only the probe call fails
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=4, sample_size=4)]
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    assert task.diagnostics["batch_mode"] is True
    assert fake_server.state.single_calls == 0
    assert task.n_scored == 4


def test_summary_matrix_is_dense_across_datasets(fake_server, write_run_config, fake_dataset):
    """One column per model when a run has a single prompt variant.

    Datasets bind different templates, so keying the headline matrix on template
    id would leave a mostly-empty grid; it is keyed on the prompt variant.
    """
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("a", n=4, sample_size=4), fake_dataset("b", n=4, sample_size=4)],
    )
    result, _ = _run(config_path)
    from abductionbench.core.reporting import _pivot, build_summary_frame

    pivot = _pivot(build_summary_frame(result))
    assert list(pivot.columns) == ["fake-model"]
    assert pivot.notna().all().all()  # no holes
    assert len(pivot.index) == 2


def test_summary_matrix_separates_prompt_modes(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("a", n=4, sample_size=4)],
        modes={"prompt_modes": ["io", "cot"]},
    )
    result, _ = _run(config_path)
    from abductionbench.core.reporting import _pivot, build_summary_frame

    pivot = _pivot(build_summary_frame(result))
    assert sorted(pivot.columns) == ["fake-model (cot)", "fake-model (io)"]


def test_output_budget_clamped_by_context_window_is_reported(
    fake_server, write_run_config, fake_dataset
):
    """A small context window silently shortens answers -- so it is counted.

    The model here has a 600-token window and the samples ask for 512 output
    tokens on top of a long prompt, so the engine must cut the output budget and
    say so instead of absorbing it.
    """
    models = [
        {
            "id": "tiny-window",
            "model_name": "test/model",
            "endpoint": {
                "base_url": fake_server.base_url,
                "api_key": "test-key",
                "batch": {"enabled": True, "group_size": 4},
            },
            "sampling": {"max_tokens_default": 512, "max_tokens_cap": 512, "max_tokens_floor": 8},
            "limits": {"max_parallel_batches": 1, "context_window": 600},
        }
    ]
    config_path = write_run_config(
        base_url=fake_server.base_url,
        models=models,
        datasets=[
            fake_dataset("fake", n=4, sample_size=4, max_tokens=512, long_every=1, long_chars=300)
        ],
        engine={
            "limits": {"input_token_budget": 4000},
            # A fine grid, so the clamped budget still leaves room to answer:
            # with the default 512-token quantum this window would leave less
            # than one quantum and the prompts would be skipped instead.
            "batching": {"max_tokens_quantum": 64},
        },
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    assert task.diagnostics["output_budgets_clamped"] == 4
    assert task.metrics["output_budget_clamped_rate"] == 1.0
    assert task.n_scored == 4  # still evaluated, just with a smaller budget


def test_no_clamp_reported_when_the_window_is_ample(fake_server, write_run_config, fake_dataset):
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=4, sample_size=4)]
    )
    result, _ = _run(config_path)
    assert result.tasks[0].metrics["output_budget_clamped_rate"] == 0.0




def test_reports_carry_the_mode_columns_and_mode_sheets(
    fake_server, write_run_config, fake_dataset
):
    """The columns item 8-10 ask for have to survive into the workbook."""
    import pandas as pd

    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=4, sample_size=4)],
        modes={"prompt_modes": ["io", "cot"]},
    )
    result, _ = _run(config_path)
    write_reports(result)
    workbook = result.run_dir / "reports" / "abductionbench_results.xlsx"
    sheets = pd.read_excel(workbook, sheet_name=None)

    assert {"Skipped_modes", "Introduced_modes"} <= set(sheets)
    for name in ("Summary_Long", "Metrics", "Tasks"):
        columns = set(sheets[name].columns)
        assert {"prompt_mode", "selection_mode", "data_delivery_mode",
                "task_kind", "template_mode"} <= columns, name

    samples = sheets["S_fake"]
    assert {"prompt_mode", "selection_mode", "data_delivery_mode", "task_kind",
            "template_mode", "response"} <= set(samples.columns)
    # Both modes are present, and each row's template_mode is its own identity.
    assert set(samples["prompt_mode"]) == {"io", "cot"}
    assert set(samples["template_mode"]) == {
        "io|n/a|generation|static",
        "cot|n/a|generation|static",
    }


def test_long_responses_are_not_truncated_to_a_preview(
    fake_server, write_run_config, fake_dataset
):
    """Item 11: the workbook holds the whole answer, not the first 2,000 chars."""
    import pandas as pd

    fake_server.state.responder = lambda conv, max_tokens: "x" * 9000
    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("fake", n=2, sample_size=2)]
    )
    result, _ = _run(config_path)
    write_reports(result)
    samples = pd.read_excel(
        result.run_dir / "reports" / "abductionbench_results.xlsx", sheet_name="S_fake"
    )
    longest = max(len(str(value)) for value in samples["response"])
    assert longest >= 9000


def test_a_prompt_with_no_room_to_answer_is_skipped_not_sent(
    fake_server, write_run_config, fake_dataset
):
    """A prompt that fills the window is this model's problem, not a failure.

    Sending it earns an HTTP 400 that takes down the whole batch it travelled
    in, so the engine skips it and says which model could not fit it.
    """
    models = [
        {
            "id": "tiny-window",
            "model_name": "test/model",
            "endpoint": {
                "base_url": fake_server.base_url,
                "api_key": "test-key",
                "batch": {"enabled": True, "group_size": 4},
            },
            "sampling": {"max_tokens_default": 4096, "max_tokens_cap": 4096},
            "limits": {"max_parallel_batches": 1, "context_window": 600},
        }
    ]
    config_path = write_run_config(
        base_url=fake_server.base_url,
        models=models,
        datasets=[
            fake_dataset("fake", n=4, sample_size=4, long_every=1, long_chars=300)
        ],
        engine={"limits": {"input_token_budget": 4000}},
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    # They count as planned -- the task set out to evaluate them -- so coverage
    # reports the loss instead of hiding it behind an empty denominator.
    assert task.n_planned == 4
    assert task.n_scored == 0
    assert task.n_skipped == 4
    assert task.metrics["coverage"] == 0.0
    # Nothing was sent, so nothing could fail.
    assert task.n_error == 0


def test_every_sheet_has_a_bold_centred_frozen_header(
    fake_server, write_run_config, fake_dataset
):
    """Checked by reading the workbook back, not by trusting the writer."""
    from openpyxl import load_workbook

    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=4, sample_size=4)],
    )
    result, _ = _run(config_path)
    write_reports(result)
    workbook = load_workbook(
        result.run_dir / "reports" / "abductionbench_results.xlsx"
    )
    assert len(workbook.sheetnames) > 5

    for name in workbook.sheetnames:
        sheet = workbook[name]
        # Row 1 is frozen everywhere. Sheets whose first column carries the row
        # labels freeze that too, so a wide matrix survives sideways scrolling.
        assert sheet.freeze_panes in ("A2", "B2"), f"{name}: {sheet.freeze_panes}"

        header = [cell for cell in sheet[1] if cell.value not in (None, "")]
        assert header, f"{name} has no header row"
        for cell in header:
            assert cell.font.bold, f"{name}!{cell.coordinate} is not bold"
            assert cell.alignment.horizontal == "center", (
                f"{name}!{cell.coordinate} is not centred"
            )


def test_repeats_multiply_the_calls_and_report_their_spread(
    fake_server, write_run_config, fake_dataset
):
    """Five calls per record, five scores, and a number saying how they differed."""
    import itertools

    # Alternate right and wrong so the repeats of a record genuinely disagree.
    answers = itertools.cycle(["right", "wrong"])
    fake_server.state.responder = lambda conv, mt: f"Answer: {next(answers)}"

    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=4, sample_size=4)],
        modes={"repeats": 5, "repeat_temperature": 0.7},
    )
    result, _ = _run(config_path)
    task = result.tasks[0]

    # Four records asked five times: twenty calls, twenty scores. Nothing folded.
    assert task.n_planned == 20
    assert task.n_scored == 20
    assert task.metrics["repeats"] == 5.0
    # The repeats disagreed, and the report says so rather than hiding it in a mean.
    assert task.metrics["repeat_agreement"] < 1.0

    records = _records(task.output_dir)
    assert len(records) == 20
    assert {r["metadata"]["repeat_of"] for r in records} == {f"s{i:04d}" for i in range(4)}
    assert {r["metadata"]["repeat_index"] for r in records} == {0, 1, 2, 3, 4}
    # Warm, and unseeded: five identical answers would measure nothing.
    assert all(r["sampling"]["temperature"] == 0.7 for r in records)
    assert all("seed" not in r["sampling"] for r in records)


def test_the_sample_sheet_carries_the_record_and_repeat_columns(
    fake_server, write_run_config, fake_dataset
):
    import pandas as pd

    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=3, sample_size=3)],
        modes={"repeats": 4},
    )
    result, _ = _run(config_path)
    write_reports(result)
    sheet = pd.read_excel(
        result.run_dir / "reports" / "abductionbench_results.xlsx", sheet_name="S_fake"
    )
    assert {"record_id", "repeat_index"} <= set(sheet.columns)
    assert len(sheet) == 12
    # Every record appears once per repeat, so the sheet can be pivoted on it.
    assert sheet.groupby("record_id")["repeat_index"].nunique().tolist() == [4, 4, 4]


def test_self_consistency_is_computed_from_the_repeats_not_bought_again(
    fake_server, write_run_config, fake_dataset
):
    """A vote over k samples needs no calls beyond the k samples.

    Three of every five answers to a record are correct, so the per-answer
    accuracy is 0.6 while the majority answer is right every time. Both numbers
    come out of the same twenty calls.
    """
    import collections
    import re as _re

    seen: collections.Counter = collections.Counter()

    def responder(conversation, max_tokens):
        body = conversation[-1]["content"]
        # The record is identified by its observation number; count how often
        # this record has been asked so the first three answers are right.
        match = _re.search(r"observation number (\d+)", body)
        key = match.group(1) if match else body[:40]
        seen[key] += 1
        if seen[key] <= 3:
            return f"Answer: echo:{body[:80]}"
        return "Answer: echo: nothing useful here"

    fake_server.state.responder = responder

    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=4, sample_size=4)],
        modes={"repeats": 5},
    )
    result, _ = _run(config_path)
    task = result.tasks[0]

    # Twenty calls for four records, in five batches of four: no extra call was
    # made for the vote.
    assert task.n_scored == 20
    assert task.checkpoint.batches_submitted == 5

    assert task.metrics["repeats"] == 5.0
    assert task.metrics["self_consistency_n_records"] == 4.0
    # Three of five answers right per record: the average answer scores 0.6,
    # the majority answer scores 1.0.
    assert task.metrics["accuracy"] == pytest.approx(0.6)
    assert task.metrics["self_consistency_accuracy"] == 1.0
    # And the spread is reported, so a flat mean cannot hide the disagreement.
    assert task.metrics["repeat_agreement"] == 0.0
    assert task.metrics["accuracy_repeat_std"] > 0.0


def test_a_vote_no_repeat_could_parse_stays_a_failure(
    fake_server, write_run_config, fake_dataset
):
    """The vote must not invent an answer none of the samples produced."""
    fake_server.state.responder = lambda conv, mt: ""

    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("fake", n=2, sample_size=2)],
        modes={"repeats": 3},
    )
    result, _ = _run(config_path)
    task = result.tasks[0]
    assert task.metrics.get("self_consistency_accuracy", 0.0) == 0.0


def test_a_judged_dataset_gets_no_vote_only_a_spread(
    fake_server, write_run_config, fake_dataset, tmp_path, monkeypatch
):
    """A plurality needs answers that can coincide.

    Free-text hypotheses graded by overlap never repeat verbatim, so every
    sample would be its own plurality of one and the "voted" score would be
    whichever sample came first. Those datasets report the spread instead.
    """
    # A unique module name: tests/judged_adapter.py already exists, and a
    # tmp_path copy of that name would shadow it for every later test.
    adapter_src = tmp_path / "overlap_scored_adapter.py"
    adapter_src.write_text(
        '''
from fake_adapter import FakeAdapter


class OverlapScoredAdapter(FakeAdapter):
    """Scored by overlap with a reference, the way a judged dataset is."""

    primary_metric = "rouge_l"
    objective_metrics = False

    def score(self, sample, response, *, output_contract=None):
        from abductionbench.core.metrics import rouge_l
        from abductionbench.core.types import SampleScore

        text = response.text
        return SampleScore(
            metrics={"rouge_l": rouge_l(text, str(sample.reference))["f"]},
            prediction=text,
        )

    def aggregate(self, scores):
        values = [s.metrics["rouge_l"] for s in scores]
        return {"rouge_l": sum(values) / max(1, len(values))}
''',
        encoding="utf-8",
    )
    monkeypatch.syspath_prepend(str(tmp_path))

    counter = iter(range(1000))
    fake_server.state.responder = lambda conv, mt: f"Answer: explanation {next(counter)}"

    config_path = write_run_config(
        base_url=fake_server.base_url,
        datasets=[
            {
                "id": "judged",
                "impl": "overlap_scored_adapter:OverlapScoredAdapter",
                "sample_size": 3,
                "options": {"n": 3},
            }
        ],
        modes={"repeats": 4},
    )
    result, _ = _run(config_path)
    task = result.tasks[0]

    assert task.n_scored == 12               # the repeats still happened
    assert task.metrics["repeats"] == 4.0
    assert "rouge_l_repeat_std" in task.metrics        # the spread is reported
    # ... but nothing was voted on.
    assert not any(k.startswith("self_consistency_") for k in task.metrics)


# --------------------------------------------------------------------------- #
# continuing a run: by id, with new work added, and from a backup
# --------------------------------------------------------------------------- #


def test_continuing_a_run_reuses_answers_and_runs_only_what_is_new(
    fake_server, write_run_config, fake_dataset
):
    """The case that matters: stop, add a dataset, continue.

    The first dataset's answers must be reused rather than paid for again, and
    the dataset that did not exist before must run in full.
    """
    first = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("a", n=4, sample_size=4)]
    )
    result, _ = _run(first)
    assert result.tasks[0].n_scored == 4
    calls_after_first = fake_server.state.requests

    # A dataset is added, and the run is continued into the same directory.
    second = write_run_config(
        base_url=fake_server.base_url,
        datasets=[
            fake_dataset("a", n=4, sample_size=4),
            fake_dataset("b", n=4, sample_size=4),
        ],
        name="test-run-2",
    )
    config = load_run_config(second)
    engine = EvaluationEngine(config, run_id=result.run_id, run_dir=result.run_dir)
    continued = asyncio.run(engine.run())

    by_dataset = {task.identity.dataset_id: task for task in continued.tasks}
    assert by_dataset["a"].n_reused == 4      # nothing re-asked
    assert by_dataset["a"].n_scored == 4      # but still scored and reported
    assert by_dataset["b"].n_reused == 0      # the new dataset ran in full
    assert by_dataset["b"].n_scored == 4
    # Only the new dataset cost anything.
    assert fake_server.state.requests - calls_after_first <= 2


def test_continuing_with_a_new_prompt_mode_runs_only_that_mode(
    fake_server, write_run_config, fake_dataset
):
    first = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("a", n=4, sample_size=4)],
        modes={"prompt_modes": ["io"]},
    )
    result, _ = _run(first)

    second = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("a", n=4, sample_size=4)],
        modes={"prompt_modes": ["io", "cot"]},
        name="test-run-2",
    )
    config = load_run_config(second)
    engine = EvaluationEngine(config, run_id=result.run_id, run_dir=result.run_dir)
    continued = asyncio.run(engine.run())

    by_mode = {task.identity.prompt_mode: task for task in continued.tasks}
    assert by_mode["io"].n_reused == 4        # the io answers stand
    assert by_mode["cot"].n_reused == 0       # cot is new work
    assert by_mode["cot"].n_scored == 4


def test_continuing_keeps_the_config_each_pass_actually_ran(
    fake_server, write_run_config, fake_dataset
):
    """Otherwise the records on disk are explained by a config that never ran."""
    first = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("a", n=2, sample_size=2)]
    )
    result, _ = _run(first)
    assert (result.run_dir / "run_config.resolved.yaml").exists()

    second = write_run_config(
        base_url=fake_server.base_url,
        datasets=[fake_dataset("a", n=2, sample_size=2), fake_dataset("b", n=2, sample_size=2)],
        name="test-run-2",
    )
    config = load_run_config(second)
    engine = EvaluationEngine(config, run_id=result.run_id, run_dir=result.run_dir)
    asyncio.run(engine.run())

    # The original is untouched and the second pass is recorded beside it.
    assert (result.run_dir / "run_config.resolved.yaml").exists()
    assert (result.run_dir / "run_config.resolved.2.yaml").exists()
    assert not (result.run_dir / ".run_config.current.yaml").exists()


def test_a_run_id_resolves_to_its_directory_under_the_output_root(
    write_run_config, fake_dataset, fake_server, tmp_path
):
    """A person has the folder name, not a path -- that is what Drive shows."""
    from abductionbench.cli import _resolve_resume

    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("a", n=2, sample_size=2)]
    )
    result, _ = _run(config_path)
    config = load_run_config(config_path)

    # By bare id...
    assert _resolve_resume(config, result.run_id) == result.run_dir
    # ... and by path, for a run kept somewhere else.
    assert _resolve_resume(config, str(result.run_dir)) == result.run_dir


def test_an_unknown_run_id_says_what_is_available_instead_of_guessing(
    write_run_config, fake_dataset, fake_server
):
    import typer

    config_path = write_run_config(
        base_url=fake_server.base_url, datasets=[fake_dataset("a", n=2, sample_size=2)]
    )
    result, _ = _run(config_path)
    config = load_run_config(config_path)

    from abductionbench.cli import _resolve_resume

    with pytest.raises(typer.Exit):
        _resolve_resume(config, "20200101-000000_not-a-run")
    # The real run is still there, untouched, and no empty directory was left.
    assert result.run_dir.exists()
    assert not (result.run_dir.parent / "20200101-000000_not-a-run").exists()
